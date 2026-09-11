import unittest
import pandas as pd
from io import BytesIO
import openpyxl
from docx import Document

from src.reporting.word_export import generate_word_report, WordExportEngine

class TestWordExportFidelity(unittest.TestCase):
    def test_flujo_word_export_matches_excel_template(self):
        # Create a mock openpyxl workbook representing an Excel template
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Flujo"
        
        ws.cell(row=8, column=1, value="Descripción")
        ws.cell(row=8, column=2, value="Nota")
        ws.cell(row=8, column=3, value="2026-07")
        ws.cell(row=8, column=4, value="2025-12")

        # Row 1: Section Header (Bold in Excel)
        ws.cell(row=9, column=1, value="Flujos de efectivo procedentes de (utilizados en) actividades de operación")
        ws.cell(row=9, column=1).font = openpyxl.styles.Font(bold=True)

        # Row 2: Data item (Regular)
        ws.cell(row=10, column=1, value="Cobros procedentes de las ventas de bienes y prestación de servicios")
        ws.cell(row=10, column=3, value=1966337)
        ws.cell(row=10, column=4, value=-98847)

        # Row 3: Subtotal (Bold, with top and bottom thin borders)
        ws.cell(row=11, column=1, value="Flujos de efectivo procedentes de actividades de operación")
        ws.cell(row=11, column=1).font = openpyxl.styles.Font(bold=True)
        ws.cell(row=11, column=3, value=16164335)
        ws.cell(row=11, column=3).font = openpyxl.styles.Font(bold=True)
        ws.cell(row=11, column=3).border = openpyxl.styles.Border(
            top=openpyxl.styles.Side(style='thin'),
            bottom=openpyxl.styles.Side(style='thin')
        )

        # Row 4: Incremento row (Regular in Excel, no borders)
        ws.cell(row=12, column=1, value="Incremento (decremento) neto en efectivo y equivalentes al efectivo")
        ws.cell(row=12, column=3, value=2720738)
        ws.cell(row=12, column=4, value=-1088400)

        excel_io = BytesIO()
        wb.save(excel_io)
        excel_bytes = excel_io.getvalue()

        df = pd.DataFrame({
            "Descripción": [
                "Flujos de efectivo procedentes de (utilizados en) actividades de operación",
                "Cobros procedentes de las ventas de bienes y prestación de servicios",
                "Flujos de efectivo procedentes de actividades de operación",
                "Incremento (decremento) neto en efectivo y equivalentes al efectivo"
            ],
            "Nota": ["", "", "", ""],
            "2026-07": [None, 1966337, 16164335, 2720738],
            "2025-12": [None, -98847, None, -1088400]
        })

        # Generate Word with excel_bytes
        docx_bytes = generate_word_report(
            df,
            title="Estado de Flujos de Efectivo",
            subtitle="Expresado en M$",
            excel_bytes=excel_bytes
        )

        self.assertIsNotNone(docx_bytes)
        doc = Document(BytesIO(docx_bytes))
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]

        # Check Header: Row 0 must have white bg, bold black text
        hdr_c0 = table.rows[0].cells[0]
        self.assertTrue(hdr_c0.paragraphs[0].runs[0].font.bold)

        # Check Section Header: Row 1 must be bold (inherited from Excel)
        r1_c0 = table.rows[1].cells[0]
        self.assertTrue(r1_c0.paragraphs[0].runs[0].font.bold)

        # Check Subtotal: Row 3 must be bold (inherited from Excel)
        r3_c0 = table.rows[3].cells[0]
        self.assertTrue(r3_c0.paragraphs[0].runs[0].font.bold)

        # Check Incremento: Row 4 must NOT be bold (matches screen & Excel!)
        r4_c0 = table.rows[4].cells[0]
        self.assertFalse(r4_c0.paragraphs[0].runs[0].font.bold)

    def test_balance_word_export_compact(self):
        # Create a 40-row dataframe representing a full classified balance
        rows = []
        for i in range(40):
            rows.append({
                "Clasificación": f"Cuenta de Balance {i+1}",
                "Nota": str(i % 10 + 1) if i % 4 == 0 else "",
                "2026-07": 1000000 + i * 50000,
                "2025-12": 950000 + i * 45000
            })
        df_bal = pd.DataFrame(rows)

        docx_bal = WordExportEngine.generate_classified_balance_word(
            df=df_bal,
            title="Estado de Situación Financiera Clasificado",
            unit="M$",
            entity_name="PACIFICO CABLE SPA"
        )
        self.assertIsNotNone(docx_bal)
        doc = Document(docx_bal)
        self.assertEqual(len(doc.tables), 1)
        # Verify 41 rows (1 header + 40 data)
        self.assertEqual(len(doc.tables[0].rows), 41)

if __name__ == '__main__':
    unittest.main()
