import os
import sys
import tempfile
import unittest
import openpyxl

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from src.reporting.notes import (
    NOTE_REGISTRY,
    is_technical_sheet,
    discover_dynamic_notes,
    get_full_note_registry,
    get_notes_by_category
)
from src.reporting.note_generator import NoteGenerator


class TestDynamicNotesDiscovery(unittest.TestCase):
    """Pruebas de detección dinámica y no invasiva de pestañas de notas adicionales."""

    def test_standard_note_registry_integrity(self):
        """Verifica que el catálogo oficial NOTE_REGISTRY (#N04-#N26) se mantenga intacto."""
        self.assertIn("#N04", NOTE_REGISTRY)
        self.assertEqual(NOTE_REGISTRY["#N04"]["sheets"], ["Efectivo"])
        self.assertEqual(NOTE_REGISTRY["#N04"]["category"], "activos_corrientes")

        self.assertIn("#N26", NOTE_REGISTRY)
        self.assertEqual(NOTE_REGISTRY["#N26"]["sheets"], ["Segmentos"])
        self.assertEqual(NOTE_REGISTRY["#N26"]["category"], "resultados")
        self.assertTrue(NOTE_REGISTRY["#N26"].get("consolidated_only"))

        # Las notas estándar oficiales son 23 (#N04 a #N26)
        self.assertEqual(len(NOTE_REGISTRY), 23)

    def test_technical_sheet_filter(self):
        """Verifica que las hojas de cálculo internas/técnicas sean ignoradas."""
        self.assertTrue(is_technical_sheet("DB_DATA"))
        self.assertTrue(is_technical_sheet("Ajustes Manuales"))
        self.assertTrue(is_technical_sheet("Ajuste Manual Noviembre 2023"))
        self.assertTrue(is_technical_sheet("Calculo RLI"))
        self.assertTrue(is_technical_sheet("BCE final 2022"))
        self.assertTrue(is_technical_sheet("PPA v3_04 Abril"))
        self.assertTrue(is_technical_sheet("_temp_calc"))

        # Hojas reales no técnicas
        self.assertFalse(is_technical_sheet("Cuadros adicionales"))
        self.assertFalse(is_technical_sheet("Nota Cambios contables"))
        self.assertFalse(is_technical_sheet("Nota reexpresion"))
        self.assertFalse(is_technical_sheet("Contingencias"))

    def test_discovery_on_current_master_template(self):
        """Verifica el descubrimiento dinámico en la Plantilla de notas_v1.xlsx actual."""
        template_path = os.path.join(ROOT_DIR, "Plantilla de notas_v1.xlsx")
        if not os.path.exists(template_path):
            self.skipTest("Plantilla de notas_v1.xlsx no encontrada")

        dynamic = discover_dynamic_notes(template_path)
        self.assertIsInstance(dynamic, dict)

        # Debe haber detectado 'Cuadros adicionales'
        detected_sheets = []
        for info in dynamic.values():
            detected_sheets.extend(info.get("sheets", []))

        self.assertIn("Cuadros adicionales", detected_sheets)

        # 'Cuadros adicionales' contiene #N27.1, por lo que debe haberse asignado a #N27
        self.assertIn("#N27", dynamic)
        self.assertEqual(dynamic["#N27"]["sheets"], ["Cuadros adicionales"])
        self.assertEqual(dynamic["#N27"]["category"], "adicionales")
        self.assertTrue(dynamic["#N27"]["is_dynamic"])

    def test_full_registry_merging(self):
        """Verifica que get_full_note_registry una NOTE_REGISTRY con las dinámicas sin alterarlas."""
        full_reg = get_full_note_registry()
        # Debe contener todas las estándar
        for code in NOTE_REGISTRY:
            self.assertIn(code, full_reg)
            self.assertEqual(full_reg[code]["sheets"], NOTE_REGISTRY[code]["sheets"])

        # Y además contener las dinámicas descubiertas
        self.assertIn("#N27", full_reg)
        self.assertGreater(len(full_reg), len(NOTE_REGISTRY))

    def test_get_notes_by_category_includes_adicionales(self):
        """Verifica que la agrupación por categorías incluya la categoría 'adicionales' ordenada."""
        cats = get_notes_by_category()
        self.assertIn("adicionales", cats)
        adic_list = cats["adicionales"]
        self.assertGreater(len(adic_list), 0)

        # Verificar que los elementos sean tuplas (código, etiqueta)
        codes = [item[0] for item in adic_list]
        self.assertIn("#N27", codes)

        # Verificar orden correlativo
        code_nums = [int(c.replace("#N", "")) for c in codes if c.startswith("#N")]
        self.assertEqual(code_nums, sorted(code_nums))

    def test_simulation_adding_new_custom_tab(self):
        """
        Simula a un usuario agregando una pestaña nueva 'Nota Contingencias Judiciales'
        al Excel y comprueba que el sistema la detecte al instante sin tocar código.
        """
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            wb = openpyxl.Workbook()
            # Hoja estándar
            ws_ef = wb.active
            ws_ef.title = "Efectivo"
            ws_ef.cell(row=1, column=1, value="Saldo Efectivo")

            # Hoja técnica a ignorar
            ws_tech = wb.create_sheet(title="Ajustes Manuales")
            ws_tech.cell(row=1, column=1, value="Ajuste X")

            # NUEVA PESTAÑA PERSONALIZADA AGREGADA POR EL USUARIO
            ws_new = wb.create_sheet(title="Nota Contingencias Judiciales")
            ws_new.cell(row=1, column=1, value="Detalle de juicios y demandas en curso")
            ws_new.cell(row=2, column=1, value="Causa A-1234")
            ws_new.cell(row=2, column=2, value=5000000)

            wb.save(tmp_path)
            wb.close()

            # Descubrir sobre esta plantilla
            discovered = discover_dynamic_notes(tmp_path)

            # Debe haber descubierto exactamente la nueva pestaña
            self.assertEqual(len(discovered), 1)
            discovered_code = list(discovered.keys())[0]
            info = discovered[discovered_code]

            self.assertEqual(info["title"], "Nota Contingencias Judiciales")
            self.assertEqual(info["sheets"], ["Nota Contingencias Judiciales"])
            self.assertEqual(info["category"], "adicionales")
            self.assertTrue(info["is_dynamic"])

            # Comprobar que en get_full_note_registry se integre perfectamente
            full_reg = get_full_note_registry(tmp_path)
            self.assertIn(discovered_code, full_reg)

            # Comprobar que aparezca en get_notes_by_category
            cats = get_notes_by_category(tmp_path)
            adic_codes = [c[0] for c in cats["adicionales"]]
            self.assertIn(discovered_code, adic_codes)

        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)


if __name__ == "__main__":
    unittest.main()
