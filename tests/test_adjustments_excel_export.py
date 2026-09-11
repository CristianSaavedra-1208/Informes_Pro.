import unittest
import io
import openpyxl
import pandas as pd
from sqlalchemy import func

from src.models.database import SessionLocal
from src.models.consolidacion import ConsolidationGroup, ConsolidationJournalEntry
from src.reporting.adjustments_export import get_adjustments_dataset, generate_adjustments_excel

class TestAdjustmentsExcelExport(unittest.TestCase):
    def setUp(self):
        self.db = SessionLocal()
        self.grupo = self.db.query(ConsolidationGroup).first()
        if not self.grupo:
            # Si no existe grupo en base, crear uno de prueba
            self.grupo = ConsolidationGroup(
                nombre_grupo="Grupo Test Auditoría",
                empresa_matriz="Db Terra Chile Holdco SpA",
                empresa_filial="Pacifico Cable SpA"
            )
            self.db.add(self.grupo)
            self.db.commit()
            self.db.refresh(self.grupo)

    def tearDown(self):
        self.db.close()

    def test_01_db_inmutability_during_export(self):
        """
        Verifica que la exportación sea 100% de solo lectura y no altere ningún registro
        ni saldo en la base de datos antes y después de su ejecución.
        """
        grupo_id = self.grupo.id

        # Capturar baseline en BD
        count_before = self.db.query(func.count(ConsolidationJournalEntry.id)).filter_by(grupo_id=grupo_id).scalar() or 0
        debe_before = self.db.query(func.sum(ConsolidationJournalEntry.debe)).filter_by(grupo_id=grupo_id).scalar() or 0.0
        haber_before = self.db.query(func.sum(ConsolidationJournalEntry.haber)).filter_by(grupo_id=grupo_id).scalar() or 0.0

        # Ejecutar todas las modalidades de exportación
        _ = get_adjustments_dataset(grupo_id, filter_mode="ALL", db=self.db)
        _ = generate_adjustments_excel(grupo_id, filter_mode="ALL", db=self.db)
        _ = generate_adjustments_excel(grupo_id, filter_mode="SINGLE_MONTH", periodo="2026-05", db=self.db)
        _ = generate_adjustments_excel(grupo_id, filter_mode="RANGE", periodo_inicio="2025-01", periodo_fin="2026-12", db=self.db)
        _ = generate_adjustments_excel(grupo_id, filter_mode="SINGLE_VOUCHER", voucher_codigo="AST-202605-001", db=self.db)

        # Capturar estado posterior en BD
        count_after = self.db.query(func.count(ConsolidationJournalEntry.id)).filter_by(grupo_id=grupo_id).scalar() or 0
        debe_after = self.db.query(func.sum(ConsolidationJournalEntry.debe)).filter_by(grupo_id=grupo_id).scalar() or 0.0
        haber_after = self.db.query(func.sum(ConsolidationJournalEntry.haber)).filter_by(grupo_id=grupo_id).scalar() or 0.0

        self.assertEqual(count_before, count_after, "❌ La cantidad de asientos en la BD cambió tras la exportación.")
        self.assertAlmostEqual(debe_before, debe_after, places=2, msg="❌ La suma total del Debe en BD varió.")
        self.assertAlmostEqual(haber_before, haber_after, places=2, msg="❌ La suma total del Haber en BD varió.")

    def test_02_export_all_periods(self):
        """
        Verifica la exportación del historial completo de ajustes del grupo.
        """
        grupo_id = self.grupo.id
        excel_bytes = generate_adjustments_excel(grupo_id, filter_mode="ALL", db=self.db)
        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 1000, "❌ El archivo Excel generado está vacío o corrupto.")

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertIn("Libro de Ajustes", wb.sheetnames)
        self.assertIn("Resumen Comprobantes", wb.sheetnames)

        ws_det = wb["Libro de Ajustes"]
        self.assertIn("INFORMES PRO", str(ws_det["A1"].value))
        self.assertIn("LIBRO DIARIO DE AJUSTES", str(ws_det["A2"].value))

    def test_03_export_single_month(self):
        """
        Verifica que al filtrar por un mes específico solo se extraigan asientos de ese período.
        """
        grupo_id = self.grupo.id
        # Obtener un período existente con asientos si hay
        sample_entry = self.db.query(ConsolidationJournalEntry).filter_by(grupo_id=grupo_id).first()
        periodo_test = sample_entry.periodo if sample_entry else "2026-05"

        dataset = get_adjustments_dataset(grupo_id, filter_mode="SINGLE_MONTH", periodo=periodo_test, db=self.db)
        for linea in dataset["lineas"]:
            self.assertEqual(linea["Período"], periodo_test, f"❌ Se encontró una línea con período distinto a {periodo_test}")

        excel_bytes = generate_adjustments_excel(grupo_id, filter_mode="SINGLE_MONTH", periodo=periodo_test, db=self.db)
        self.assertGreater(len(excel_bytes), 1000)

    def test_04_export_range_periods(self):
        """
        Verifica la exportación en un rango cronológico de períodos (ej. 2025-01 a 2026-06).
        """
        grupo_id = self.grupo.id
        p_ini = "2025-01"
        p_fin = "2026-06"

        dataset = get_adjustments_dataset(
            grupo_id,
            filter_mode="RANGE",
            periodo_inicio=p_ini,
            periodo_fin=p_fin,
            db=self.db
        )

        for linea in dataset["lineas"]:
            self.assertTrue(p_ini <= linea["Período"] <= p_fin, f"❌ La línea {linea['Código Folio']} está fuera del rango {p_ini}-{p_fin}")

        excel_bytes = generate_adjustments_excel(
            grupo_id,
            filter_mode="RANGE",
            periodo_inicio=p_ini,
            periodo_fin=p_fin,
            db=self.db
        )
        self.assertGreater(len(excel_bytes), 1000)

    def test_05_export_single_voucher(self):
        """
        Verifica la exportación de un comprobante específico por su código de folio.
        """
        grupo_id = self.grupo.id
        sample_entry = self.db.query(ConsolidationJournalEntry).filter_by(grupo_id=grupo_id).filter(ConsolidationJournalEntry.asiento_codigo != None).first()

        if sample_entry and sample_entry.asiento_codigo:
            codigo = sample_entry.asiento_codigo
            dataset = get_adjustments_dataset(grupo_id, filter_mode="SINGLE_VOUCHER", voucher_codigo=codigo, db=self.db)
            self.assertGreater(len(dataset["lineas"]), 0)
            for linea in dataset["lineas"]:
                self.assertEqual(linea["Código Folio"], codigo)

            excel_bytes = generate_adjustments_excel(grupo_id, filter_mode="SINGLE_VOUCHER", voucher_codigo=codigo, db=self.db)
            self.assertGreater(len(excel_bytes), 1000)

            wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
            self.assertIn("Libro de Ajustes", wb.sheetnames)

if __name__ == '__main__':
    unittest.main()
