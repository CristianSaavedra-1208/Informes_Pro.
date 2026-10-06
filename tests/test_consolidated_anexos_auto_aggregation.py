import os
import sys
import unittest
import openpyxl
from io import BytesIO

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from src.core.external_notes_manager import ExternalNotesManager


class TestConsolidatedAnexosAutoAggregation(unittest.TestCase):
    """
    Pruebas para la suma y agregación automática de anexos externos en grupos consolidados,
    con exclusión estricta de Impuestos Diferidos y Saldos Intercompañía / Relacionadas.
    """

    def test_group_detection_and_subsidiary_resolution(self):
        """Verifica la detección de grupos consolidados y resolución de filiales en la base de datos."""
        self.assertTrue(ExternalNotesManager.is_consolidated_group("[GRUPO] Consolidado DB Terra Holdco"))
        self.assertTrue(ExternalNotesManager.is_consolidated_group("[GRUPO] Consolidado Terra Parent"))
        self.assertFalse(ExternalNotesManager.is_consolidated_group("Db Terra Chile Holdco SpA"))
        self.assertFalse(ExternalNotesManager.is_consolidated_group("Pacifico Cable SpA"))

        # Grupo DB Terra Holdco
        subs_holdco = ExternalNotesManager.get_group_companies("[GRUPO] Consolidado DB Terra Holdco")
        self.assertIn("Db Terra Chile Holdco SpA", subs_holdco)
        self.assertIn("Pacifico Cable SpA", subs_holdco)

        # Grupo Terra Parent (jerárquico / recursivo)
        subs_parent = ExternalNotesManager.get_group_companies("[GRUPO] Consolidado Terra Parent")
        self.assertIn("Db Terra Chile Parent SpA", subs_parent)
        self.assertIn("Db Terra Chile Holdco SpA", subs_parent)
        self.assertIn("Pacifico Cable SpA", subs_parent)

    def test_auto_aggregation_math_and_exclusions(self):
        """
        Simula dos filiales con anexos externos cargados y verifica:
        1. Suma aritmética celda por celda de notas operacionales (ej: Deudores).
        2. Exclusión estricta de 'Impuestos Diferidos'.
        3. Exclusión estricta de 'Empresas relacionadas' e 'Inversion en relacionadas'.
        """
        test_period = "2099-12"
        sub1 = "Db Terra Chile Holdco SpA"
        sub2 = "Pacifico Cable SpA"
        group_name = "[GRUPO] Consolidado DB Terra Holdco"

        master_path = ExternalNotesManager.get_master_template_path()
        if not os.path.exists(master_path):
            self.skipTest(f"Master template no encontrada: {master_path}")

        # Crear anexo para Filial 1
        wb1 = openpyxl.load_workbook(master_path)
        if "Deudores" in wb1.sheetnames:
            wb1["Deudores"].cell(7, 2, value=1500)
            wb1["Deudores"].cell(8, 2, value=500)
        if "Impuestos Diferidos" in wb1.sheetnames:
            wb1["Impuestos Diferidos"].cell(10, 2, value=999999)
        if "Empresas relacionadas" in wb1.sheetnames:
            wb1["Empresas relacionadas"].cell(10, 2, value=888888)

        buf1 = BytesIO()
        wb1.save(buf1)
        wb1.close()

        # Crear anexo para Filial 2
        wb2 = openpyxl.load_workbook(master_path)
        if "Deudores" in wb2.sheetnames:
            wb2["Deudores"].cell(7, 2, value=2500)
            wb2["Deudores"].cell(8, 2, value=1500)
        if "Impuestos Diferidos" in wb2.sheetnames:
            wb2["Impuestos Diferidos"].cell(10, 2, value=777777)
        if "Empresas relacionadas" in wb2.sheetnames:
            wb2["Empresas relacionadas"].cell(10, 2, value=666666)

        buf2 = BytesIO()
        wb2.save(buf2)
        wb2.close()

        # Guardar en rutas de filiales para el período de prueba
        path1 = ExternalNotesManager.save_anexo(sub1, test_period, buf1.getvalue())
        path2 = ExternalNotesManager.save_anexo(sub2, test_period, buf2.getvalue())

        try:
            # Verificar info del grupo consolidado (debe reportar auto-agregación activa)
            info = ExternalNotesManager.get_anexo_info(group_name, test_period)
            self.assertTrue(info["exists"])
            self.assertTrue(info["is_auto_aggregated"])
            self.assertIn(sub1, info["subsidiaries"])
            self.assertIn(sub2, info["subsidiaries"])

            # Ejecutar agregación automática
            agg_stream = ExternalNotesManager.build_aggregated_group_anexo(group_name, test_period)
            self.assertIsNotNone(agg_stream)

            agg_wb = openpyxl.load_workbook(agg_stream, data_only=False)

            # 1. Verificar SUMA AUTOMÁTICA en Deudores
            ws_deudores = agg_wb["Deudores"]
            val_r7 = ws_deudores.cell(7, 2).value
            val_r8 = ws_deudores.cell(8, 2).value
            self.assertEqual(val_r7, 4000, f"Deudores R7 debería ser 1500 + 2500 = 4000, pero dio {val_r7}")
            self.assertEqual(val_r8, 2000, f"Deudores R8 debería ser 500 + 1500 = 2000, pero dio {val_r8}")

            # 2. Verificar EXCLUSIÓN de Impuestos Diferidos
            if "Impuestos Diferidos" in agg_wb.sheetnames:
                ws_id = agg_wb["Impuestos Diferidos"]
                val_id = ws_id.cell(10, 2).value
                self.assertNotEqual(val_id, 999999 + 777777, "Impuestos Diferidos NO debe sumar saldos de filiales")

            # 3. Verificar EXCLUSIÓN de Empresas Relacionadas (Saldos intercompañía)
            if "Empresas relacionadas" in agg_wb.sheetnames:
                ws_rel = agg_wb["Empresas relacionadas"]
                val_rel = ws_rel.cell(10, 2).value
                self.assertNotEqual(val_rel, 888888 + 666666, "Empresas Relacionadas NO debe sumar saldos intercompañía")

            agg_wb.close()

        finally:
            # Limpiar archivos de prueba
            for p in [path1, path2]:
                if os.path.exists(p):
                    try:
                        os.remove(p)
                    except OSError:
                        pass


if __name__ == "__main__":
    unittest.main()
