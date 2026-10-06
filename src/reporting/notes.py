class NotesOrchestrator:
    def __init__(self, mapping_rules):
        """
        Req 5: Orquesta las tablas de Notas.
        mapping_rules define qué cuentas o categorías entran en qué Nota.
        """
        self.mapping_rules = mapping_rules
        
    def generate_note_tables(self, mapped_tb_df):
        """
        Retorna un diccionario con detalles estructurados validos para tablas Word.
        Ej: {'Nota_10_Inventarios': {'total': 5000, 'detalle': [...]}}
        """
        notes_data = {}
        for note_name, categories in self.mapping_rules.items():
            # Filtrar Trial Balance
            note_df = mapped_tb_df[mapped_tb_df['categoria_financiera'].isin(categories)].copy()
            
            if not note_df.empty:
                # Agrupar para generar la tabla descriptiva de la nota
                table_data = note_df.groupby(['cuenta_id', 'descripcion'])['saldo_final'].sum().reset_index()
                notes_data[note_name] = {
                    "total": table_data['saldo_final'].sum(),
                    "detalle": table_data.to_dict(orient='records')
                }
            else:
                notes_data[note_name] = {"total": 0.0, "detalle": []}
                
        return notes_data


# Registro Maestro de Códigos Únicos de Notas con prefijo '#'
NOTE_REGISTRY = {
    "#N04": {
        "title": "Efectivo y equivalentes al efectivo",
        "sheets": ["Efectivo"],
        "category": "activos_corrientes"
    },
    "#N05": {
        "title": "Otros activos no financieros (Corriente)",
        "sheets": ["Otros activos no financieros, c"],
        "category": "activos_corrientes"
    },
    "#N06": {
        "title": "Deudores comerciales y otras cuentas por cobrar",
        "sheets": ["Deudores"],
        "category": "activos_corrientes"
    },
    "#N07": {
        "title": "Inventarios",
        "sheets": ["Inventarios"],
        "category": "activos_corrientes"
    },
    "#N08": {
        "title": "Activos intangibles distintos de la plusvalía",
        "sheets": ["Intangibles"],
        "category": "activos_no_corrientes"
    },
    "#N09": {
        "title": "Propiedades, planta y equipo",
        "sheets": ["Activo Fijo"],
        "category": "activos_no_corrientes"
    },
    "#N10": {
        "title": "Activos por derechos de uso",
        "sheets": ["Activo por derechos de uso"],
        "category": "activos_no_corrientes"
    },
    "#N11": {
        "title": "Plusvalía",
        "sheets": ["Plusvalia"],
        "category": "activos_no_corrientes"
    },
    "#N12": {
        "title": "Activos y pasivos de impuestos corrientes",
        "sheets": ["Impuestos corrientes"],
        "category": "activos_corrientes"
    },
    "#N13": {
        "title": "Impuestos diferidos",
        "sheets": ["Impuestos Diferidos"],
        "category": "activos_no_corrientes"
    },
    "#N14": {
        "title": "Cuentas por cobrar/pagar a entidades relacionadas e inversiones",
        "sheets": ["Empresas relacionadas", "Inversion en relacionadas"],
        "category": "pasivos_corrientes"
    },
    "#N15": {
        "title": "Instrumentos / Pasivos financieros",
        "sheets": ["Pasivos financieros "],
        "category": "pasivos_no_corrientes"
    },
    "#N16": {
        "title": "Pasivos por derechos de uso",
        "sheets": ["Pasivos derechos de  uso"],
        "category": "pasivos_no_corrientes"
    },
    "#N17": {
        "title": "Cuentas por pagar comerciales y otras cuentas por pagar",
        "sheets": ["Cuentas por pagar"],
        "category": "pasivos_corrientes"
    },
    "#N18": {
        "title": "Beneficios a los empleados",
        "sheets": ["Provisiones"],  # Mapeada a la pestaña de Provisiones
        "category": "pasivos_no_corrientes"
    },
    "#N19": {
        "title": "Otros pasivos no financieros (Corriente)",
        "sheets": ["Otros pasivos no financieros"],
        "category": "pasivos_corrientes"
    },
    "#N20": {
        "title": "Patrimonio",
        "sheets": ["Patrimonio"],
        "category": "patrimonio"
    },
    "#N21": {
        "title": "Ingresos de actividades ordinarias y costo de ventas",
        "sheets": ["Ingresos Ctos operacion"],
        "category": "resultados"
    },
    "#N22": {
        "title": "Gastos de administración",
        "sheets": ["Gtos Adm"],
        "category": "resultados"
    },
    "#N23": {
        "title": "Diferencia de cambio",
        "sheets": ["DC y Reajustes"],
        "category": "resultados"
    },
    "#N24": {
        "title": "Costos e ingresos financieros",
        "sheets": ["Costos e ingresos Financieros"],
        "category": "resultados"
    },
    "#N25": {
        "title": "Otros ingresos y egresos",
        "sheets": ["Otros gastos por funcion", "Otros ingresos por funcion"],
        "category": "resultados"
    },
    "#N26": {
        "title": "Segmentos de operación",
        "sheets": ["Segmentos"],
        "category": "resultados",
        "consolidated_only": True
    }
}

import os
import re

TECHNICAL_SHEET_KEYWORDS = {'db_data', 'ajuste', 'ajustes', 'calculo', 'bce ', 'ppa '}

def is_technical_sheet(sheet_name: str) -> bool:
    """Identifica hojas auxiliares o técnicas de cálculo que no son notas para presentación."""
    s = str(sheet_name).strip().lower()
    if s.startswith('_') or s.startswith('temp'):
        return True
    for kw in TECHNICAL_SHEET_KEYWORDS:
        if kw in s:
            return True
    return False

# Caché en memoria para evitar reabrir el Excel en cada interacción si el archivo no cambió
_DYNAMIC_NOTES_CACHE = {}

def discover_dynamic_notes(template_path: str = "Plantilla de notas_v1.xlsx") -> dict:
    """
    Inspecciona la plantilla Excel y detecta automáticamente pestañas adicionales
    que no forman parte de NOTE_REGISTRY ni sean hojas técnicas.
    
    Asigna códigos explícitos si los encuentra en la hoja (ej. #N27.1 -> #N27)
    o códigos correlativos (#N28, #N29, etc.).
    """
    if not template_path or not os.path.exists(template_path):
        fallback = "Plantilla de notas_v1.xlsx"
        if os.path.exists(fallback):
            template_path = fallback
        else:
            return {}

    try:
        mtime = os.path.getmtime(template_path)
    except OSError:
        mtime = 0

    cache_key = (os.path.abspath(template_path), mtime)
    if cache_key in _DYNAMIC_NOTES_CACHE:
        return _DYNAMIC_NOTES_CACHE[cache_key]

    import openpyxl

    known_sheets = set()
    max_n = 26
    for code, info in NOTE_REGISTRY.items():
        m = re.findall(r'#N(\d+)', code)
        if m:
            max_n = max(max_n, int(m[0]))
        for sh in info.get('sheets', []):
            known_sheets.add(sh.strip().lower())

    try:
        wb = openpyxl.load_workbook(template_path, data_only=True, read_only=True)
        sheet_names = list(wb.sheetnames)
        wb.close()
    except Exception:
        return {}

    candidate_sheets = [
        s for s in sheet_names 
        if s.strip().lower() not in known_sheets and not is_technical_sheet(s)
    ]

    if not candidate_sheets:
        _DYNAMIC_NOTES_CACHE[cache_key] = {}
        return {}

    # Abrir para escanear cabeceras si hay códigos explícitos
    dynamic_notes = {}
    try:
        wb = openpyxl.load_workbook(template_path, data_only=True)
        assigned_codes = set(NOTE_REGISTRY.keys())
        
        # 1er paso: Buscar si alguna hoja tiene códigos explícitos (#Nxx)
        explicit_sheet_codes = {}
        for sh in candidate_sheets:
            if sh not in wb.sheetnames:
                continue
            ws = wb[sh]
            found_nums = []
            for r in range(1, min(25, ws.max_row + 1)):
                for c in range(1, min(10, ws.max_column + 1)):
                    v = str(ws.cell(r, c).value or '')
                    m = re.findall(r'#N(\d+)', v)
                    if m:
                        found_nums.extend([int(x) for x in m])
            if found_nums:
                target_num = min(found_nums)
                cand_code = f"#N{target_num:02d}"
                if cand_code not in assigned_codes:
                    explicit_sheet_codes[sh] = cand_code
                    assigned_codes.add(cand_code)
                    if target_num > max_n:
                        max_n = target_num

        # 2do paso: Asignar códigos correlativos para las demás hojas adicionales
        for sh in candidate_sheets:
            if sh in explicit_sheet_codes:
                code = explicit_sheet_codes[sh]
            else:
                max_n += 1
                code = f"#N{max_n:02d}"
                assigned_codes.add(code)

            dynamic_notes[code] = {
                "title": sh.strip(),
                "sheets": [sh],
                "category": "adicionales",
                "is_dynamic": True
            }

        wb.close()
    except Exception:
        # Fallback simple sin inspección profunda
        for idx, sh in enumerate(candidate_sheets, start=max_n + 1):
            code = f"#N{idx:02d}"
            dynamic_notes[code] = {
                "title": sh.strip(),
                "sheets": [sh],
                "category": "adicionales",
                "is_dynamic": True
            }

    _DYNAMIC_NOTES_CACHE[cache_key] = dynamic_notes
    return dynamic_notes


def get_full_note_registry(template_path: str = "Plantilla de notas_v1.xlsx") -> dict:
    """
    Retorna el registro completo unificado: NOTE_REGISTRY oficial + Notas dinámicas detectadas.
    """
    dynamic = discover_dynamic_notes(template_path)
    if not dynamic:
        return dict(NOTE_REGISTRY)
    merged = dict(NOTE_REGISTRY)
    merged.update(dynamic)
    return merged


def get_notes_by_category(template_path: str = "Plantilla de notas_v1.xlsx") -> dict:
    """
    Retorna la estructura de notas agrupadas por categoría, incluyendo la categoría 'adicionales'.
    """
    categories = {
        "activos_corrientes": [],
        "activos_no_corrientes": [],
        "pasivos_corrientes": [],
        "pasivos_no_corrientes": [],
        "patrimonio": [],
        "resultados": [],
        "adicionales": []
    }
    
    full_registry = get_full_note_registry(template_path)
    for code, info in full_registry.items():
        cat = info.get('category', 'adicionales')
        if cat not in categories:
            categories[cat] = []
        categories[cat].append((code, f"[{code}] {info['title']}"))

    def _sort_key(item):
        m = re.findall(r'(\d+)', item[0])
        return int(m[0]) if m else 999

    for cat in categories:
        categories[cat].sort(key=_sort_key)

    return categories

