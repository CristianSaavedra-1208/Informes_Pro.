import streamlit as st
import pandas as pd
import os
import datetime
from src.core.excel_utils import df_to_excel_bytes, format_periodo, read_excel_cached

from src.reporting.notes import NOTE_REGISTRY, get_full_note_registry, get_notes_by_category

# Mapeo de notas dinámico reconstruido desde NOTE_REGISTRY y plantillas activas
sheet_name_map = {code: info['sheets'] for code, info in get_full_note_registry().items()}

# Agrupación de notas dinámica agrupada por categorías del estado financiero
notes_by_category = get_notes_by_category()

def load_all_entity_contexts(active_entity, periodo_actual, periodo_comp, map_balance_df, map_pl_df):
    from src.core.sabana_manager import SabanaManager
    bundle = SabanaManager.get_sabana_bundle(active_entity, periodo_actual, periodo_comp, map_balance_df, map_pl_df)
    return bundle['contexts']


import datetime

def evaluate_openpyxl_formula(formula, ws, col_idx, row_idx=None, visited=None):
    import re
    from openpyxl.utils import column_index_from_string
    if not isinstance(formula, str) or not formula.startswith('='):
        return formula
        
    if visited is None:
        visited = set()
        
    if row_idx is not None:
        coord = (row_idx, col_idx)
        if coord in visited:
            return 0.0
        visited.add(coord)
        
    clean_formula = formula.strip().upper()
    
    # 1. Handle functions like SUM(...) or SUMA(...)
    # Soportar =SUM(A1:A5), =SUMA(C5:C11), =SUM(C1:C3, C5:C6), =SUMA(C1; C2; C3)
    def resolve_cell_or_range(term):
        term = term.strip()
        if not term:
            return 0.0
        if ':' in term:
            parts = term.split(':')
            m1 = re.match(r'^([A-Z]+)(\d+)$', parts[0].strip())
            m2 = re.match(r'^([A-Z]+)(\d+)$', parts[1].strip())
            if m1 and m2:
                c1_idx = column_index_from_string(m1.group(1))
                r1_num = int(m1.group(2))
                c2_idx = column_index_from_string(m2.group(1))
                r2_num = int(m2.group(2))
                subtotal = 0.0
                min_r, max_r = min(r1_num, r2_num), max(r1_num, r2_num)
                min_c, max_c = min(c1_idx, c2_idx), max(c1_idx, c2_idx)
                for r in range(min_r, min_r + (max_r - min_r + 1)):
                    if r <= ws.max_row:
                        for c in range(min_c, min_c + (max_c - min_c + 1)):
                            if c <= ws.max_column:
                                val = ws.cell(row=r, column=c).value
                                if isinstance(val, str) and val.startswith('='):
                                    val = evaluate_openpyxl_formula(val, ws, c, r, visited)
                                try:
                                    if val is not None:
                                        subtotal += float(val)
                                except (ValueError, TypeError):
                                    pass
                return subtotal
        else:
            m = re.match(r'^([A-Z]+)(\d+)$', term)
            if m:
                c_idx = column_index_from_string(m.group(1))
                r_num = int(m.group(2))
                if r_num <= ws.max_row and c_idx <= ws.max_column:
                    val = ws.cell(row=r_num, column=c_idx).value
                    if isinstance(val, str) and val.startswith('='):
                        val = evaluate_openpyxl_formula(val, ws, c_idx, r_num, visited)
                    try:
                        if val is not None:
                            return float(val)
                    except (ValueError, TypeError):
                        pass
        try:
            return float(term)
        except (ValueError, TypeError):
            return 0.0

    # Match SUM / SUMA with arguments inside
    func_match = re.match(r'^=(?:SUM|SUMA)\((.*)\)$', clean_formula)
    if func_match:
        inner = func_match.group(1)
        # Dividir argumentos por coma o punto y coma
        args = [arg.strip() for arg in re.split(r'[,;]', inner) if arg.strip()]
        total = sum(resolve_cell_or_range(arg) for arg in args)
        if row_idx is not None:
            visited.discard((row_idx, col_idx))
        return total

    # 2. Arithmetic expressions: e.g. =+C26+C18+C12, =C32+C30+C28, =+D34, =C10-C20
    # Remover el signo '=' inicial
    expr = clean_formula[1:].strip()
    if expr.startswith('+'):
        expr = expr[1:].strip()

    # Encontrar todas las celdas o rangos referenciados en la expresión
    tokens = re.findall(r'([A-Z]+)(\d+)', expr)
    if tokens:
        # Reemplazar tokens de mayor longitud a menor longitud para evitar solapamientos
        sorted_tokens = sorted(set(tokens), key=lambda t: len(t[0] + t[1]), reverse=True)
        for col_letter, row_str in sorted_tokens:
            token_str = f"{col_letter}{row_str}"
            r_num = int(row_str)
            c_idx = column_index_from_string(col_letter)
            val_float = 0.0
            if r_num <= ws.max_row and c_idx <= ws.max_column:
                val = ws.cell(row=r_num, column=c_idx).value
                if isinstance(val, str) and val.startswith('='):
                    val = evaluate_openpyxl_formula(val, ws, c_idx, r_num, visited)
                if val is not None:
                    try:
                        val_float = float(val)
                    except (ValueError, TypeError):
                        val_float = 0.0
            # Reemplazar usando regex de límite de palabra/carácter
            expr = re.sub(r'\b' + token_str + r'\b', f"({val_float})", expr)

        # Reemplazar operadores duplicados
        expr = expr.replace('+-', '-').replace('--', '+').replace('++', '+')
        if re.match(r'^[0-9. +\-*/()]+$', expr):
            try:
                ret_val = eval(expr)
                if row_idx is not None:
                    visited.discard((row_idx, col_idx))
                return float(ret_val)
            except Exception:
                pass

    if row_idx is not None:
        visited.discard((row_idx, col_idx))
    return None

def evaluate_formulas_in_workbook(excel_bytes_in):
    import openpyxl
    from io import BytesIO
    excel_bytes_in.seek(0)
    wb = openpyxl.load_workbook(excel_bytes_in, data_only=False)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                val = cell.value
                if isinstance(val, str) and val.startswith('='):
                    visited = set()
                    evaluated = evaluate_openpyxl_formula(val, ws, col, row, visited)
                    if evaluated is not None:
                        cell.value = evaluated
    
    excel_bytes_out = BytesIO()
    wb.save(excel_bytes_out)
    excel_bytes_out.seek(0)
    return excel_bytes_out

def classify_row(row):
    cells = [val for val in row if pd.notna(val) and str(val).strip() != ""]
    if len(cells) == 0:
        return "empty", None
    
    # Check if it contains any cell starting with "validacion" (case-insensitive)
    is_validation = False
    for cell in row:
        if pd.notna(cell) and str(cell).strip().lower().startswith("validacion"):
            is_validation = True
            break
            
    if is_validation:
        return "validation", cells
        
    if len(cells) == 1:
        label = str(cells[0]).strip().lower()
        if any(k in label for k in ["total", "saldo final", "sub-total", "subtotal", "totales"]):
            return "table", cells
        return "text", cells[0]
        
    return "table", cells

def is_header_start_row(row):
    # Contar celdas no vacías para omitir filas de detalle con fechas
    non_empty_cells = [val for val in row if pd.notna(val) and str(val).strip() != ""]
    if len(non_empty_cells) > 3:
        if any(isinstance(val, datetime.datetime) for val in row):
            return False
        row_strs = [str(val).strip().lower() for val in row if pd.notna(val)]
        if any(any(k in s for k in ["clp", "us$", "uf", "préstamo", "leasing", "tasa", "banco", "prestamo"]) for s in row_strs):
            return False

    has_date = False
    has_entity = False
    has_numeric_value = False
    
    for val in row:
        if pd.isna(val) or val == "":
            continue
        if isinstance(val, (int, float)):
            if val > 2000 and val < 2100:
                has_date = True
            else:
                has_numeric_value = True
        elif isinstance(val, datetime.datetime):
            has_date = True
        else:
            val_str = str(val).strip().lower()
            if any(k in val_str for k in ["31.12.", "31.03.", "30.06.", "30.09.", "31 de ", "30 de "]):
                has_date = True
            if any(k in val_str for k in ["pacifico", "holdco", "consolidado", "matriz", "filial", "empresa"]):
                has_entity = True
                
    return (has_date or has_entity) and not has_numeric_value

def get_base_title(text):
    import re
    if not text:
        return ""
    s = re.sub(r'\[COMPARATIVO\]', '', str(text), flags=re.IGNORECASE).strip().lower()
    s = re.sub(r'^\d+[\)\.\-]\s*', '', s).strip()
    s = re.sub(r'^[a-z][\)\.\-]\s*', '', s).strip()
    s = re.sub(r'\b(202\d|201\d|199\d)\b', '', s).strip()
    s = re.sub(r'\b\d{2}[\.\/]\d{2}[\.\/]\d{2,4}\b', '', s).strip()
    return " ".join(s.split())

def is_same_base_title(t1, t2):
    b1 = get_base_title(t1)
    b2 = get_base_title(t2)
    if not b1 or not b2:
        return False
    return b1 == b2 or b1 in b2 or b2 in b1

def is_main_section_title(first_cell_txt, full_row_title, non_empty):
    import re
    first_low = first_cell_txt.lower().strip()
    full_low = full_row_title.lower().strip()
    
    # Must contain at least some alphabetic characters (not just dates or numbers)
    if not re.search(r'[a-zA-ZáéíóúÁÉÍÓÚñÑ]', full_row_title):
        return False
        
    # Ignore table header labels, column names, or total rows
    if first_low in ['total', 'totales', 'saldo final', 'subtotal', 'sub-total', 'concepto', 'detalle', 'activos', 'pasivos', 'institución', 'institucion']:
        return False
    if first_low.startswith('total') or first_low.startswith('saldo final') or first_low.startswith('sub total') or first_low.startswith('sub-total'):
        return False
        
    numeric_vals = [v for v in non_empty if isinstance(v, (int, float)) and not (1990 <= v <= 2030)]
    if numeric_vals:
        return False
        
    # Numbered titles: e.g. 1), 2), 3), 4), 5), a), b), 1.-
    if re.match(r'^\d+[\)\.\-]', first_cell_txt) or re.match(r'^[a-z][\)\.\-]', first_cell_txt):
        return True
        
    explicit_starters = [
        'vencimiento de pasivos', 'vencimientos', 'cambios en pasivo', 
        'detalle leasing', 'detalle prestamos', 'detalle prstamos', 'detalle de instrumentos',
        'detalle de fondos mutuos', 'reconciliacion de tasa', 'reconciliacin de tasa', 'reconciliación de tasa',
        'composicion del gasto', 'composicin del gasto', 'composición del gasto',
        'los saldos acumulados', 'conciliacion del saldo', 'conciliacin del saldo', 'conciliación del saldo',
        'los efectos de impuestos', 'nota de ', 'provisiones por pasivos', 'contratos con clientes',
        'movimientos de propiedades', 'movimientos de activos', 'detalle de leasing', 'detalle de prestamos'
    ]
    
    if any(full_low.startswith(starter) for starter in explicit_starters):
        return True
        
    return False

def split_sheet_into_elements(df, sheet_name=None):
    import re
    df = df.dropna(how='all', axis=1)
    
    code_pattern = re.compile(r'#([A-Za-z0-9_]+(?:\.[A-Za-z0-9_]+)?)')
    
    # 1. PRE-ESCANEO: ¿Contiene la hoja códigos explícitos con '#'? (ej: #N15.1, #N09.2, [#N15.2])
    explicit_codes_found = []
    for idx, row in df.iterrows():
        for val in row.values:
            if pd.notna(val) and isinstance(val, str):
                m = code_pattern.search(val.strip())
                if m:
                    explicit_codes_found.append(f"#{m.group(1)}")
                    
    is_explicit_mode = len(explicit_codes_found) > 0

    # =========================================================================
    # MODO EXPLÍCITO: La plantilla Excel gobierna estrictamente con '#'
    # =========================================================================
    if is_explicit_mode:
        elements = []
        current_chunk = []
        current_explicit_code = None
        has_data = False
        in_table = False
        
        for idx, row in df.iterrows():
            row_vals = list(row.values)
            non_empty = [v for v in row_vals if pd.notna(v) and str(v).strip() != ""]
            if not non_empty:
                if current_chunk and in_table and not has_data:
                    current_chunk = []
                elif current_chunk and has_data:
                    current_chunk.append(row_vals)
                continue
                
            first_cell_txt = str(non_empty[0]).strip()
            full_row_title = " ".join([str(v).strip() for v in non_empty if isinstance(v, str)])
            
            m = code_pattern.search(first_cell_txt) or code_pattern.search(full_row_title)
            
            if m:
                found_code = f"#{m.group(1)}"
                
                # Si es el mismo código actualmente activo (ej: réplica duplicada del período comparativo)
                if current_explicit_code is not None and found_code == current_explicit_code:
                    continue
                    
                # Si es un nuevo código explícito, guardar la tabla previa
                if current_chunk and has_data:
                    chunk_df = pd.DataFrame(current_chunk).dropna(how='all', axis=0)
                    if not chunk_df.empty:
                        elements.append(("table", chunk_df, sheet_name, current_explicit_code))
                    current_chunk = []
                    has_data = False
                    in_table = False
                    
                current_explicit_code = found_code
                
                # Extraer título descriptivo limpio (sin códigos técnicos ni [COMPARATIVO])
                clean_title = re.sub(r'\[?#' + re.escape(m.group(1)) + r'\]?', '', full_row_title)
                clean_title = re.sub(r'\[COMPARATIVO\]', '', clean_title, flags=re.IGNORECASE).strip()
                clean_title = re.sub(r'\b(202\d|201\d|199\d)\b', '', clean_title).strip()
                clean_title = re.sub(r'^[\s\-:]+', '', clean_title).strip()
                if clean_title:
                    elements.append(("text", clean_title, sheet_name, current_explicit_code))
                continue
            else:
                has_data = True
                in_table = True
                current_chunk.append(row_vals)
                
        if current_chunk and has_data:
            chunk_df = pd.DataFrame(current_chunk).dropna(how='all', axis=0)
            if not chunk_df.empty:
                elements.append(("table", chunk_df, sheet_name, current_explicit_code))
                
        return elements

    # =========================================================================
    # MODO LEGADO DE RESPALDO: Para notas que aún no tienen '#' en su plantilla
    # =========================================================================
    elements = []
    current_chunk = []
    has_data = False
    in_table = False

    active_pasivos_section = None

    for idx, row in df.iterrows():
        row_vals = list(row.values)
        non_empty = [v for v in row_vals if pd.notna(v) and str(v).strip() != ""]
        if not non_empty:
            if current_chunk and in_table and not has_data:
                current_chunk = []
            elif current_chunk and has_data:
                current_chunk.append(row_vals)
            continue
            
        first_cell_txt = str(non_empty[0]).strip()
        full_row_title = " ".join([str(v).strip() for v in non_empty if isinstance(v, str)])
        
        is_split = is_main_section_title(first_cell_txt, full_row_title, non_empty)
        
        # Regla exclusiva para Activo Fijo: dividir en Cuadro #N09.1 (Detalle por rubro) y Cuadro #N09.2 (Cuadro de movimientos comparativo)
        if not is_split and sheet_name and str(sheet_name).strip() == "Activo Fijo":
            if current_chunk and has_data:
                if len(non_empty) == 1 and re.match(r'^20\d{2}$', first_cell_txt):
                    # Solo dividir la primera vez (para separar #N09.1 de #N09.2)
                    if len([e for e in elements if e[0] == "table"]) == 0:
                        is_split = True

        # Regla exclusiva para Pasivos Financieros (#N15):
        # Mantener los dos cuadros comparativos de "vencimiento de pasivos" y "cambios en pasivo" bajo un solo código
        if sheet_name and "pasivos financieros" in str(sheet_name).strip().lower():
            full_low = full_row_title.lower()
            if is_split:
                if "vencimiento de pasivos" in full_low:
                    if active_pasivos_section == "vencimientos":
                        is_split = False
                        continue
                    else:
                        active_pasivos_section = "vencimientos"
                elif "cambios en pasivo" in full_low:
                    if active_pasivos_section == "cambios":
                        is_split = False
                        continue
                    else:
                        active_pasivos_section = "cambios"
                else:
                    active_pasivos_section = "other"

        if is_split:
            if current_chunk and has_data:
                chunk_df = pd.DataFrame(current_chunk).dropna(how='all', axis=0)
                if not chunk_df.empty:
                    elements.append(("table", chunk_df))
                current_chunk = []
                has_data = False
                in_table = False
                
            if is_main_section_title(first_cell_txt, full_row_title, non_empty):
                clean_txt = re.sub(r'\[COMPARATIVO\]', '', full_row_title, flags=re.IGNORECASE).strip()
                clean_header = re.sub(r'\b(202\d|201\d|199\d)\b', '', clean_txt).strip()
                if clean_header.strip():
                    elements.append(("text", clean_header.strip()))
            else:
                has_data = True
                in_table = True
                current_chunk.append(row_vals)
        else:
            has_data = True
            in_table = True
            current_chunk.append(row_vals)

    if current_chunk and has_data:
        chunk_df = pd.DataFrame(current_chunk).dropna(how='all', axis=0)
        if not chunk_df.empty:
            elements.append(("table", chunk_df))
        
    return elements

def render_note_section(notes_list, key_prefix, scale_factor_nota, unidad_nota, empresa_path, empresa_seleccionada, periodo_actual, periodo_comp):
    options = [n[1] for n in notes_list]
    selected_label = st.selectbox(
        "Seleccione la nota a generar:",
        options,
        key=f"{key_prefix}_select"
    )
    
    # Buscar nota seleccionada
    selected_id = None
    for n_id, n_label in notes_list:
        if n_label == selected_label:
            selected_id = n_id
            break
            
    if selected_id:
        st.write("---")
        
        # Validar si la nota está en la plantilla y tiene pestañas asignadas
        if selected_id not in sheet_name_map or not sheet_name_map[selected_id]:
            st.warning(f"⚠️ La nota '{selected_label}' no tiene pestañas Excel configuradas en la plantilla (es de carácter informativo o no se encuentra implementada aún).")
            return
            
        col_b1, col_b2 = st.columns([2, 2])
        with col_b1:
            run_note = st.button("🚀 Ejecutar y Visualizar Nota", type="primary", use_container_width=True, key=f"{key_prefix}_run")
        with col_b2:
            if st.button("🔄 Recargar Saldos y Asientos", type="secondary", use_container_width=True, key=f"{key_prefix}_reload"):
                from src.core.sabana_manager import SabanaManager
                SabanaManager.clear_sabana_cache()
                st.success("🔄 Caché de saldos y asientos actualizada. Haz clic en '🚀 Ejecutar y Visualizar Nota'.")
            
        target_lang = st.session_state.get('idioma_reporte', 'es')
        cache_key = f"_note_result__{key_prefix}__{selected_id}__{target_lang}"

        if run_note:
            from src.core.sabana_manager import SabanaManager
            SabanaManager.clear_sabana_cache()

            template_nota = "Plantilla de notas_v1.xlsx"
            if not os.path.exists(template_nota):
                st.error("❌ No se encontró la plantilla maestra de notas 'Plantilla de notas_v1.xlsx' en la raíz.")
            else:
                st.info(f"Procesando y mapeando saldos para {selected_label}...")
                
                # Cargar mapeos con caché (evita re-lectura si el archivo no cambió)
                map_bal_local = os.path.join(empresa_path, "map_balance.xlsx")
                map_pl_local  = os.path.join(empresa_path, "map_pl.xlsx")
                map_balance_df = read_excel_cached(map_bal_local, dtype=str) if os.path.exists(map_bal_local) else None
                map_pl_df      = read_excel_cached(map_pl_local,  dtype=str) if os.path.exists(map_pl_local)  else None

                if map_balance_df is None:
                    st.error("❌ No se pudo cargar el maestro de Mapeo Balance.")
                    st.stop()

                import importlib
                import src.reporting.note_generator as note_gen_module
                importlib.reload(note_gen_module)
                NoteGenerator = note_gen_module.NoteGenerator

                try:
                    # Invalidar caché previa de resultados para forzar recalculo fresco
                    if cache_key in st.session_state:
                        del st.session_state[cache_key]

                    # Cargar contextos con caché en session_state.
                    _ctx_cache_key = f"_entity_ctx__{empresa_seleccionada}__{periodo_actual}__{periodo_comp}"
                    if _ctx_cache_key not in st.session_state:
                        with st.spinner("Cargando saldos y compilando contextos de datos..."):
                            st.session_state[_ctx_cache_key] = load_all_entity_contexts(
                                active_entity=empresa_seleccionada,
                                periodo_actual=periodo_actual,
                                periodo_comp=periodo_comp,
                                map_balance_df=map_balance_df,
                                map_pl_df=map_pl_df
                            )
                    entity_contexts = st.session_state[_ctx_cache_key]

                    is_consolidated = empresa_seleccionada.startswith("[GRUPO]")
                    target_sheets = sheet_name_map[selected_id]
                    
                    engine = NoteGenerator(template_nota)
                    excel_nota_out = engine.generate(
                        sheet_names=target_sheets,
                        entity_contexts=entity_contexts,
                        active_entity_name=empresa_seleccionada,
                        is_consolidated=is_consolidated,
                        scale_factor=scale_factor_nota,
                        periodo_actual_str=periodo_actual,
                        periodo_comp_str=periodo_comp,
                        map_balance_df=map_balance_df,
                        map_pl_df=map_pl_df,
                        target_lang=target_lang
                    )
                    
                    # Inyectar datos de anexos externos (si existen para la empresa/período)
                    import importlib
                    import src.core.external_notes_manager as ext_mgr_module
                    importlib.reload(ext_mgr_module)
                    ExternalNotesManager = ext_mgr_module.ExternalNotesManager
                    excel_nota_out = ExternalNotesManager.inject_external_data(
                        excel_nota_out,
                        empresa=empresa_seleccionada,
                        periodo_actual=periodo_actual,
                        periodo_comp=periodo_comp if periodo_comp != "Ninguno" else None,
                        target_lang=target_lang
                    )
                    excel_eval_out = evaluate_formulas_in_workbook(excel_nota_out)
                    preview_nota_df = pd.read_excel(excel_eval_out, sheet_name=target_sheets[0], header=None)
                    
                    # --- SEMÁFORO FINANCIERO DE CUADRATURA (TIE-OUT) ---
                    tie_out_type = None
                    tie_out_text = None
                    try:
                        from src.models.database import SessionLocal
                        from src.models.historical_data import HistoricalDataRecord
                        
                        NOTE_TO_RUBRO_MAP = {
                             "#N04": [("Balance", "Efectivo y efectivo equivalente")],
                             "#N05": [("Balance", "Otros activos no financieros, corrientes")],
                             "#N06": [("Balance", "Deudores comerciales y otras cuentas por cobrar, corrientes")],
                             "#N07": [("Balance", "Inventarios")],
                             "#N08": [("Balance", "Activos intangibles distinto a la plusvalía"), ("Balance", "Activos intangibles distinto a la plusvalia")],
                             "#N09": [("Balance", "Propiedades, plantas y equipos")],
                             "#N10": [("Balance", "Activo por derechos de uso")],
                             "#N11": [("Balance", "Plusvalia"), ("Balance", "Plusvalía")],
                             "#N12": [("Balance", "Activo por impuestos, corrientes"), ("Balance", "Pasivo por impuestos, corrientes")],
                             "#N13": [("Balance", "Activo por impuestos diferidos, no corrientes")],
                             "#N14": [("Balance", "Cuentas por cobrar a entidades relacionadas, no corrientes"),
                                      ("Balance", "Cuentas por pagar entidades relacionadas, corrientes"), 
                                      ("Balance", "Cuentas por pagar entidades relacionadas, no corrientes"),
                                      ("Balance", "Inversion en empresas relacionadas")],
                             "#N15": [("Balance", "Otros pasivos financieros corrientes"), ("Balance", "Otros pasivos financieros, no corriente")],
                             "#N16": [("Balance", "Pasivos por derechos de uso, corrientes"), ("Balance", "Pasivos por derechos de uso, no corrientes")],
                             "#N17": [("Balance", "Cuentas comerciales y otras cuentas por pagar, corrientes"),
                                      ("Balance", "Cuentas comerciales y otras cuentas por pagar, no corrientes")],
                             "#N18": [("Balance", "Provisiones por beneficios a los empleados")],
                             "#N19": [("Balance", "Otras provisiones, no corrientes")],
                             "#N20": [("Balance", "Capital emitido"), ("Balance", "Otras reservas"), ("Balance", "Resultados acumulados")],
                             "#N21": [("P&L", "Ingresos de actividades ordinarias")],
                             "#N22": [("P&L", "Gastos de administracin"), ("P&L", "Gastos de administración")],
                             "#N23": [("P&L", "Diferencias de cambio"), ("P&L", "Diferencia de cambio")],
                             "#N24": [("P&L", "Costos financieros")],
                             "#N25": [("P&L", "Otros ingresos por funcin"), ("P&L", "Otros ingresos por funcion"), ("P&L", "Otros egresos por funcin"), ("P&L", "Otros egresos por funcion")]
                        }
                        
                        rubros = NOTE_TO_RUBRO_MAP.get(selected_id)
                        if rubros:
                            db = SessionLocal()
                            db_total = 0.0
                            for rep, lin in rubros:
                                r_recs = db.query(HistoricalDataRecord).filter_by(
                                    empresa=empresa_seleccionada,
                                    periodo=periodo_actual,
                                    reporte=rep,
                                    linea_item=lin
                                ).all()
                                for r_rec in r_recs:
                                    db_total += abs(float(r_rec.monto))
                            db.close()
                            
                            # Buscar la fila de "Total" en el preview_nota_df
                            total_row_idx = None
                            for idx, row_series in preview_nota_df.iterrows():
                                row_str = " ".join([str(val).strip().lower() for val in row_series.values if pd.notna(val)])
                                if "total" in row_str and "sub" not in row_str:
                                    if any(k in row_str for k in ["corriente", "general", "administrac", "operac", "bruto", "neto", "saldo"]):
                                        total_row_idx = idx
                                        break
                                        
                            if total_row_idx is None:
                                for idx, row_series in preview_nota_df.iterrows():
                                    row_str = " ".join([str(val).strip().lower() for val in row_series.values if pd.notna(val)])
                                    if "total" in row_str and "sub" not in row_str:
                                        total_row_idx = idx
                                        break
                                    
                            if total_row_idx is not None:
                                total_row = preview_nota_df.iloc[total_row_idx]
                                numeric_cols = []
                                import datetime
                                import re
                                for col_idx, val in enumerate(total_row.values):
                                    if pd.notna(val) and isinstance(val, (int, float)):
                                        is_date_col = False
                                        for r in range(0, total_row_idx):
                                            cell_val = preview_nota_df.iloc[r, col_idx]
                                            if pd.isna(cell_val) or str(cell_val).strip() == "":
                                                for c in range(col_idx - 1, -1, -1):
                                                    left_val = preview_nota_df.iloc[r, c]
                                                    if pd.notna(left_val) and str(left_val).strip() != "":
                                                        cell_val = left_val
                                                        break
                                            if pd.notna(cell_val) and str(cell_val).strip() != "":
                                                if isinstance(cell_val, (datetime.datetime, pd.Timestamp)):
                                                    is_date_col = True
                                                    break
                                                if len(str(cell_val)) <= 25 and re.search(r'\b(202\d|199\d)\b', str(cell_val)):
                                                    is_date_col = True
                                                    break
                                        if is_date_col:
                                            numeric_cols.append((col_idx, float(val)))
                                        
                                excel_total = None
                                if numeric_cols:
                                    if len(numeric_cols) == 1:
                                        excel_total = abs(numeric_cols[0][1]) * scale_factor_nota
                                    else:
                                        neto_col_idx = None
                                        for col_idx, val in numeric_cols:
                                            for r in range(0, total_row_idx):
                                                cell_val = preview_nota_df.iloc[r, col_idx]
                                                if pd.notna(cell_val) and "neto" in str(cell_val).strip().lower():
                                                    neto_col_idx = col_idx
                                                    break
                                            if neto_col_idx is not None:
                                                break
                                        
                                        if neto_col_idx is not None:
                                            for col_idx, val in numeric_cols:
                                                if col_idx == neto_col_idx:
                                                    excel_total = abs(val) * scale_factor_nota
                                                    break
                                        else:
                                            excel_total = abs(numeric_cols[0][1]) * scale_factor_nota
                                            
                                if excel_total is not None:
                                    diff = abs(excel_total - db_total)
                                    if diff < 1000.0:
                                        tie_out_type = "success"
                                        tie_out_text = f"⚖️ **[Tie-Out OK]** El total de la Nota (${excel_total:,.0f}) coincide perfectamente con el Estado Financiero Principal (${db_total:,.0f})."
                                    else:
                                        tie_out_type = "warning"
                                        tie_out_text = f"⚖️ **[Tie-Out Diferencia]** Se detectó un descuadre de **${diff:,.0f}** entre el total de la Nota (${excel_total:,.0f}) y los Estados Financieros (${db_total:,.0f})."
                        else:
                            tie_out_type = "info"
                            tie_out_text = "ℹ️ Esta nota no tiene rubros principales en el balance asociados para validación automática."
                    except Exception as tie_err:
                        tie_out_type = "info"
                        tie_out_text = f"ℹ️ Validación de Tie-Out omitida temporalmente: {tie_err}"
                    
                    # Intentar extraer usando rangos nombrados
                    from src.core.excel_utils import extract_named_ranges_from_excel
                    named_ranges = extract_named_ranges_from_excel(excel_nota_out, selected_id)
                    
                    elements = []
                    if named_ranges:
                        for r_name, df_chunk in named_ranges:
                            elements.append(("table", df_chunk, target_sheets[0]))
                    else:
                        for sheet_name in target_sheets:
                            sheet_df = pd.read_excel(excel_eval_out, sheet_name=sheet_name, header=None)
                            if len(target_sheets) > 1:
                                elements.append(("text", f"Pestaña: {sheet_name.strip()}", sheet_name, None))
                            sheet_elements = split_sheet_into_elements(sheet_df, sheet_name=sheet_name)
                            for s_item in sheet_elements:
                                el_type = s_item[0]
                                el_val = s_item[1]
                                s_name = s_item[2] if len(s_item) > 2 else sheet_name
                                exp_code = s_item[3] if len(s_item) > 3 else None
                                elements.append((el_type, el_val, s_name, exp_code))
                    
                    from src.reporting.word_export import WordExportEngine
                    word_nota_out = WordExportEngine.generate_notes_word(
                        elements=elements,
                        title=selected_label,
                        unit=unidad_nota,
                        note_code=selected_id,
                        excel_bytes=excel_eval_out.getvalue(),
                        target_lang=target_lang
                    )
                    
                    excel_nota_out.seek(0)
                    excel_eval_out.seek(0)
                    if hasattr(word_nota_out, "seek"):
                        word_nota_out.seek(0)
                        word_bytes_val = word_nota_out.getvalue()
                    else:
                        word_bytes_val = word_nota_out

                    st.session_state[cache_key] = {
                        "selected_label": selected_label,
                        "unidad_nota": unidad_nota,
                        "selected_id": selected_id,
                        "excel_bytes": excel_nota_out.getvalue(),
                        "eval_bytes": excel_eval_out.getvalue(),
                        "word_bytes": word_bytes_val,
                        "elements": elements,
                        "target_sheets": target_sheets,
                        "tie_out_type": tie_out_type,
                        "tie_out_text": tie_out_text
                    }
                except Exception as e:
                    st.error(f"❌ Error al procesar nota: {e}")
                    st.exception(e)

        # Mostrar contenido si existe en cache de session_state
        if cache_key in st.session_state:
            cached = st.session_state[cache_key]
            st.success(f"✅ Nota generada exitosamente en {cached['unidad_nota']}.")
                
            from src.reporting.formatting import apply_corporate_style
            from io import BytesIO
            eval_io = BytesIO(cached["eval_bytes"])
            
            table_counter = 0
            for item in cached["elements"]:
                if len(item) == 4:
                    el_type, el_val, sh_name, explicit_code = item
                elif len(item) == 3:
                    el_type, el_val, sh_name = item
                    explicit_code = None
                else:
                    el_type, el_val = item[0], item[1]
                    sh_name = cached["target_sheets"][0]
                    explicit_code = None
                    
                if el_type == "text":
                    clean_text = str(el_val).replace("[COMPARATIVO]", "").replace("[comparativo]", "").strip()
                    if not clean_text:
                        continue
                    from src.core.narrative_translator import translate_narrative_text
                    display_text = translate_narrative_text(clean_text, target_lang=target_lang)
                    st.markdown(f"##### 📋 {display_text}")
                else:
                    chunk_df = el_val.dropna(how='all', axis=0).reset_index(drop=True)
                    if chunk_df.empty:
                        continue
                    
                    if explicit_code:
                        sub_code = explicit_code
                    else:
                        table_counter += 1
                        sub_code = f"{cached['selected_id']}.{table_counter}"
                    label_cuadro = "Schedule" if str(target_lang).lower() == "en" else "Cuadro"
                    st.markdown(f"**📍 {label_cuadro} `{sub_code}`**")
                    
                    new_cols = [" " * (idx + 1) for idx in range(len(chunk_df.columns))]
                    chunk_df.columns = new_cols
                    chunk_df = chunk_df.fillna("")
                    
                    styled_df = apply_corporate_style(chunk_df, excel_bytes=eval_io, sheet_name=sh_name, target_lang=target_lang)
                    html_str = styled_df.to_html(index=False)
                    import re
                    html_str = re.sub(r'<thead\b[^>]*>.*?</thead>', '', html_str, flags=re.DOTALL)
                    scrollable_html = f'<div style="overflow-x: auto; width: 100%;">{html_str}</div>'
                    st.markdown(scrollable_html, unsafe_allow_html=True)
                    st.write("")

            st.write("")
            c1, c2, c3 = st.columns([2, 2, 1])
            with c1:
                st.download_button(
                    "📥 Descargar Nota en Excel",
                    data=cached["excel_bytes"],
                    file_name=f"{cached['selected_label'].replace(' ', '_')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary",
                    use_container_width=True,
                    key=f"{key_prefix}_dl_ex_{selected_id}"
                )
            with c2:
                st.download_button(
                    "📝 Descargar Nota en Word",
                    data=cached["word_bytes"],
                    file_name=f"{cached['selected_label'].replace(' ', '_')}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    type="primary",
                    use_container_width=True,
                    key=f"{key_prefix}_dl_wd_{selected_id}"
                )
            with c3:
                def _cb_clear_cache(ck):
                    if ck in st.session_state:
                        del st.session_state[ck]
                st.button(
                    "🗑️ Limpiar Vista",
                    type="secondary",
                    use_container_width=True,
                    key=f"{key_prefix}_clear_{selected_id}",
                    on_click=_cb_clear_cache,
                    args=(cache_key,)
                )

def render(empresa_seleccionada, empresa_path):
    st.title("📑 Informes y Notas a los Estados Financieros")
    
    if "GLOBAL" in empresa_seleccionada:
        st.info("🌐 **Modo Global Activo**: Desde esta sección puedes administrar las **Plantillas Maestras del Sistema**: la Plantilla Global de Notas (Contable) y la Plantilla Base de Anexos (Extracontable para filiales). Ambas plantillas aplican como base para todas las empresas.")

        if 'global_tpl_feedback' in st.session_state:
            fb_type, fb_msg = st.session_state.pop('global_tpl_feedback')
            if fb_type == "success":
                st.success(fb_msg)
            elif fb_type == "error":
                st.error(fb_msg)

        tab_g1, tab_g2 = st.tabs([
            "📘 Plantilla Maestra Global de Notas (Contable)",
            "📙 Plantilla Base Máster de Anexos (Extracontable)"
        ])

        with tab_g1:
            st.subheader("📘 Plantilla Maestra Global de Notas Contables (`Plantilla de notas_v1.xlsx`)")
            st.write("Esta es la **base y molde maestro de todas las notas del sistema**. Si agregas nuevas notas, tablas o códigos (`#N...`), hazlo en este archivo y vuelve a subirlo.")
            
            template_nota = "Plantilla de notas_v1.xlsx"
            if os.path.exists(template_nota):
                mtime_g = int(os.path.getmtime(template_nota))
                size_g = os.path.getsize(template_nota) / 1024.0
                dt_g = datetime.datetime.fromtimestamp(mtime_g).strftime("%Y-%m-%d %H:%M:%S")
                st.caption(f"🕒 **Última modificación:** {dt_g} | 📦 **Tamaño:** {size_g:.1f} KB")

                with open(template_nota, "rb") as file:
                    fresh_bytes_g = file.read()
                st.download_button(
                    label="📥 Descargar Plantilla Maestra Global Actual (`Plantilla de notas_v1.xlsx`)",
                    data=fresh_bytes_g,
                    file_name="Plantilla_de_notas_v1.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                    key=f"download_global_master_template_notes_{mtime_g}"
                )
            else:
                st.warning("⚠️ No se encontró la plantilla maestra global 'Plantilla de notas_v1.xlsx' en la raíz.")

            st.write("")
            st.markdown("##### 📤 Subir y Actualizar Plantilla Maestra Global")
            uploaded_template = st.file_uploader(
                "Seleccionar archivo Excel para reemplazar la Plantilla Maestra Global (.xlsx):",
                type=["xlsx"],
                key="global_notes_template_uploader"
            )
            
            if uploaded_template is not None:
                if st.button("💾 Guardar y Actualizar Plantilla Maestra Global", type="primary", use_container_width=True, key="save_global_master_notes_btn"):
                    try:
                        with open(template_nota, "wb") as f:
                            f.write(uploaded_template.getbuffer())
                        # Limpiar caché de notas
                        for k in list(st.session_state.keys()):
                            if k.startswith("_note_result__") or k.startswith("_entity_ctx__"):
                                del st.session_state[k]
                        st.session_state['global_tpl_feedback'] = ("success", "✅ ¡Plantilla Maestra Global de Notas subida y actualizada con éxito en la raíz del proyecto!")
                        st.rerun()
                    except PermissionError:
                        st.error("❌ El archivo 'Plantilla de notas_v1.xlsx' está abierto en Excel. Ciérralo y vuelve a presionar Guardar.")
                    except Exception as e:
                        st.error(f"❌ Error al guardar la plantilla: {e}")

        with tab_g2:
            st.subheader("📙 Plantilla Base Máster de Anexos Extracontables (`Plantilla EXTERNA_de_notas_v1.xlsx`)")
            st.write("Esta es la **plantilla base oficial de anexos** que descargan las filiales para completar datos cualitativos y operacionales (antigüedad de deuda, vencimientos, etc.), basada en las notas de la plantilla global.")
            
            from src.core.external_notes_manager import ExternalNotesManager
            ext_info = ExternalNotesManager.get_master_template_info()
            if ext_info["exists"]:
                st.caption(f"🕒 **Última modificación:** {ext_info['last_modified']} | 📦 **Tamaño:** {ext_info['size_kb']}")
                if ext_info["sheets"]:
                    st.caption(f"📑 **{len(ext_info['sheets'])} Hojas disponibles:** {', '.join(ext_info['sheets'][:6])}{'...' if len(ext_info['sheets']) > 6 else ''}")

                with open(ext_info["path"], "rb") as mf:
                    fresh_bytes_ext = mf.read()
                st.download_button(
                    label="📥 Descargar Plantilla Base Máster de Anexos (`Plantilla EXTERNA_de_notas_v1.xlsx`)",
                    data=fresh_bytes_ext,
                    file_name="Plantilla_EXTERNA_de_notas_v1.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                    key=f"download_global_master_ext_notes_{ext_info['mtime']}"
                )
            else:
                st.warning("⚠️ No se encontró la plantilla base de anexos en la raíz.")

            st.write("")
            st.markdown("##### 📤 Subir y Actualizar Plantilla Base Máster de Anexos")
            up_master_anexo = st.file_uploader(
                "Seleccionar archivo Excel para reemplazar la Plantilla Base Máster de Anexos (.xlsx):",
                type=["xlsx"],
                key="global_master_anexo_uploader"
            )
            if up_master_anexo is not None:
                if st.button("💾 Guardar y Actualizar Plantilla Base de Anexos Global", type="primary", use_container_width=True, key="save_global_master_anexo_btn"):
                    try:
                        ExternalNotesManager.save_master_template(up_master_anexo.getbuffer())
                        # Limpiar caché de notas
                        for k in list(st.session_state.keys()):
                            if k.startswith("_note_result__") or k.startswith("_entity_ctx__"):
                                del st.session_state[k]
                        st.session_state['global_tpl_feedback'] = ("success", "✅ ¡Plantilla Base Máster de Anexos actualizada con éxito en la raíz del sistema!")
                        st.rerun()
                    except PermissionError:
                        st.error("❌ El archivo 'Plantilla EXTERNA_de_notas_v1.xlsx' está abierto en Excel. Ciérralo y vuelve a presionar Guardar.")
                    except Exception as e:
                        st.error(f"❌ Error al guardar la plantilla: {e}")

        st.stop()

        
    from src.models.trial_balance_db import TrialBalanceDB
    
    # Obtener períodos disponibles unificados (Memoria Activa + Histórico)
    periodos_hist = TrialBalanceDB.get_available_periods(empresa_seleccionada)
    if not periodos_hist:
        periodos_hist = ["2026-12", "2026-07", "2026-06", "2026-03", "2025-12"]
        
    # Parámetros globales en columnas
    col_p1, col_p2, col_esc = st.columns([1.5, 1.5, 2])
    with col_p1:
        periodo_actual = st.selectbox("Periodo Actual:", periodos_hist, index=0, format_func=format_periodo)
    with col_p2:
        periodo_comp = st.selectbox("Periodo Comparativo (Opcional):", ["Ninguno"] + periodos_hist, index=1 if len(periodos_hist) > 1 else 0, format_func=format_periodo)
    with col_esc:
        unidad_nota = st.selectbox(
            "Unidad de medida en reportes:",
            ["Miles de pesos (M$)", "Unidades ($)", "Millones de pesos (MM$)"],
            index=0
        )
        
        if "Miles" in unidad_nota:
            scale_factor_nota = 1000.0
        elif "Millones" in unidad_nota:
            scale_factor_nota = 1000000.0
        else:
            scale_factor_nota = 1.0
            
    st.write("")
    
    main_tab_viz, main_tab_anexos = st.tabs([
        "👁️ Visualizar y Exportar Notas",
        "📥 Carga de Anexos Operacionales por Empresa"
    ])

    with main_tab_viz:
        is_consolidated = empresa_seleccionada.startswith("[GRUPO] ")

        # Cargar registro completo actualizado desde la plantilla maestra global
        template_nota_activa = "Plantilla de notas_v1.xlsx"

        full_registry = get_full_note_registry(template_nota_activa)
        sheet_name_map.update({code: info['sheets'] for code, info in full_registry.items()})
        current_notes_by_cat = get_notes_by_category(template_nota_activa)

        # Filtrar notas por categoría según modo (individual vs consolidado)
        active_notes_by_cat = {}
        for cat, note_list in current_notes_by_cat.items():
            filtered_list = []
            for code, label in note_list:
                info = full_registry.get(code, {})
                if info.get("consolidated_only") and not is_consolidated:
                    continue
                filtered_list.append((code, label))
            active_notes_by_cat[cat] = filtered_list

        # Crear pestañas para cada rubro, incorporando Notas Adicionales dinámicas
        tabs = st.tabs([
            "💰 Activos Corrientes",
            "🏢 Activos No Corrientes",
            "💳 Pasivos Corrientes",
            "🛡️ Pasivos No Corrientes",
            "📊 Patrimonio",
            "📈 Resultados",
            "📌 Notas Adicionales",
            "📂 Informes Corporativos"
        ])
        
        # Tab 1: Activos Corrientes
        with tabs[0]:
            st.subheader("Notas de Activos Corrientes")
            render_note_section(active_notes_by_cat.get("activos_corrientes", []), "act_corr", scale_factor_nota, unidad_nota, empresa_path, empresa_seleccionada, periodo_actual, periodo_comp)
            
        # Tab 2: Activos No Corrientes
        with tabs[1]:
            st.subheader("Notas de Activos No Corrientes")
            render_note_section(active_notes_by_cat.get("activos_no_corrientes", []), "act_no_corr", scale_factor_nota, unidad_nota, empresa_path, empresa_seleccionada, periodo_actual, periodo_comp)
            
        # Tab 3: Pasivos Corrientes
        with tabs[2]:
            st.subheader("Notas de Pasivos Corrientes")
            render_note_section(active_notes_by_cat.get("pasivos_corrientes", []), "pas_corr", scale_factor_nota, unidad_nota, empresa_path, empresa_seleccionada, periodo_actual, periodo_comp)
            
        # Tab 4: Pasivos No Corrientes
        with tabs[3]:
            st.subheader("Notas de Pasivos No Corrientes")
            render_note_section(active_notes_by_cat.get("pasivos_no_corrientes", []), "pas_no_corr", scale_factor_nota, unidad_nota, empresa_path, empresa_seleccionada, periodo_actual, periodo_comp)
            
        # Tab 5: Patrimonio
        with tabs[4]:
            st.subheader("Notas de Patrimonio")
            render_note_section(active_notes_by_cat.get("patrimonio", []), "patrimonio", scale_factor_nota, unidad_nota, empresa_path, empresa_seleccionada, periodo_actual, periodo_comp)
            
        # Tab 6: Resultados
        with tabs[5]:
            st.subheader("Notas de Resultados (P&L)")
            render_note_section(active_notes_by_cat.get("resultados", []), "resultados", scale_factor_nota, unidad_nota, empresa_path, empresa_seleccionada, periodo_actual, periodo_comp)

        # Tab 7: Notas Adicionales / Personalizadas detectadas dinámicamente
        with tabs[6]:
            st.subheader("📌 Notas y Cuadros Adicionales (Detectados de la Plantilla)")
            adic_notes = active_notes_by_cat.get("adicionales", [])
            if adic_notes:
                st.caption("Pestañas personalizadas detectadas automáticamente en la plantilla Excel activa:")
                render_note_section(adic_notes, "adic_notes", scale_factor_nota, unidad_nota, empresa_path, empresa_seleccionada, periodo_actual, periodo_comp)
            else:
                st.info("ℹ️ No se detectaron pestañas adicionales en la plantilla actual. Si agregas una pestaña nueva a la plantilla Excel de notas, el sistema la detectará aquí automáticamente.")
            
        # Tab 8: Informes Corporativos
        with tabs[7]:
            st.subheader("Informes Adicionales")
            st.write("Generación de paquetes de gerencia corporativos. (PDF, Word).")
            st.info("Funcionalidad en desarrollo.")

    with main_tab_anexos:
        st.subheader("📥 Carga e Integración de Anexos / Notas Extracontables")
        st.write(
            "Desde esta sección puedes administrar y cargar la **Plantilla Externa de Anexos** por empresa y período "
            "para aquellas notas que contienen datos operacionales o desgloses estructurados "
            "(ej: Fondos Mutuos, Antigüedad de Deudores, Vencimiento de Pasivos, Inversiones en Relacionadas, etc.)."
        )
        
        from src.core.external_notes_manager import ExternalNotesManager
        
        col_anx1, col_anx2 = st.columns([2, 2])
        with col_anx1:
            anexo_periodo = st.selectbox(
                "Seleccionar Período de Carga:",
                periodos_hist,
                index=0,
                format_func=format_periodo,
                key="anexo_ext_period_select"
            )
        
        with col_anx2:
            info = ExternalNotesManager.get_anexo_info(empresa_seleccionada, anexo_periodo)
            if info["exists"]:
                if info.get("is_auto_aggregated"):
                    subs_str = ", ".join(info.get("subsidiaries", []))
                    st.info(
                        f"⚡ **Consolidado Automático Activo**: Este grupo consolida y suma automáticamente las filiales con anexo registrado (**{subs_str}**).\n\n"
                        f"*(Excluye automáticamente 'Impuestos Diferidos' y 'Saldos Intercompañía/Relacionadas')*"
                    )
                    agg_bytes = ExternalNotesManager.build_aggregated_group_anexo(empresa_seleccionada, anexo_periodo)
                    if agg_bytes:
                        st.download_button(
                            label="📥 Descargar Anexo Consolidado Pre-sumado (.xlsx)",
                            data=agg_bytes,
                            file_name=f"Anexo_Consolidado_PreSumado_{anexo_periodo}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True,
                            key=f"download_auto_agg_{empresa_seleccionada}_{anexo_periodo}"
                        )
                else:
                    st.success(f"✅ Anexo registrado para **{empresa_seleccionada}** ({format_periodo(anexo_periodo)})\n\n🕒 **Modificado:** {info['last_modified']} | 📦 **Tamaño:** {info['size_kb']}")
            else:
                st.warning(f"⚠️ Sin anexo externo registrado para **{empresa_seleccionada}** en **{format_periodo(anexo_periodo)}**.")

        # Mensaje de retroalimentación persistente tras guardar
        if 'anexo_feedback' in st.session_state:
            fb_type, fb_msg = st.session_state.pop('anexo_feedback')
            if fb_type == "success":
                st.success(fb_msg)
            elif fb_type == "error":
                st.error(fb_msg)
            elif fb_type == "warning":
                st.warning(fb_msg)

        st.write("---")
        c_dn1, c_dn2 = st.columns([2, 2])
        with c_dn1:
            st.markdown("#### 1. Descargar Formato Oficial de Anexos")
            st.write("Descarga el formato base vigente del sistema para ingresar los datos operacionales de esta empresa en este período.")
            
            master_info = ExternalNotesManager.get_master_template_info()
            if master_info["exists"]:
                st.caption(f"🕒 **Plantilla oficial vigente:** {master_info['last_modified']} | 📦 **Tamaño:** {master_info['size_kb']}")
                if master_info["sheets"]:
                    sheets_preview = ", ".join(master_info["sheets"][:5])
                    more_count = len(master_info["sheets"]) - 5
                    if more_count > 0:
                        sheets_preview += f" (+{more_count} más)"
                    st.caption(f"📑 **{len(master_info['sheets'])} Hojas disponibles:** {sheets_preview}")

                try:
                    with open(master_info["path"], "rb") as mf:
                        fresh_bytes = mf.read()
                    st.download_button(
                        label="📥 Descargar Formato Base de Anexos (.xlsx)",
                        data=fresh_bytes,
                        file_name="Plantilla_EXTERNA_de_notas_v1.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                        key=f"download_master_external_template_btn_{master_info['mtime']}"
                    )
                except Exception as e:
                    st.error(f"Error al preparar la descarga: {e}")
            else:
                st.warning("⚠️ No se encontró la plantilla base de anexos en el sistema.")

            st.caption("💡 *Nota: Para modificar o agregar nuevas hojas a la plantilla base del sistema, dirígete a `[GLOBAL] Configuración General` en la barra lateral izquierda.*")


        with c_dn2:
            st.markdown("#### 2. Subir Anexo Completado")
            st.write(f"Sube el archivo Excel completado para **{empresa_seleccionada}** en el período **{format_periodo(anexo_periodo)}**:")
            uploaded_anexo = st.file_uploader(
                "Seleccionar archivo Excel (.xlsx):",
                type=["xlsx"],
                key=f"anexo_uploader_{empresa_seleccionada}_{anexo_periodo}"
            )
            if uploaded_anexo is not None:
                if st.button("💾 Guardar Anexo de la Empresa para este Período", type="primary", use_container_width=True, key="save_anexo_btn"):
                    try:
                        saved_path = ExternalNotesManager.save_anexo(empresa_seleccionada, anexo_periodo, uploaded_anexo.getbuffer())
                        # Limpiar caché de notas para forzar recarga inmediata con los nuevos datos
                        for k in list(st.session_state.keys()):
                            if k.startswith("_note_result__") or k.startswith("_entity_ctx__"):
                                del st.session_state[k]
                        st.session_state['anexo_feedback'] = ("success", f"✅ ¡Anexo guardado e integrado con éxito para **{empresa_seleccionada}** ({format_periodo(anexo_periodo)})! Los datos ya están listos en la pestaña de Notas.")
                        st.rerun()
                    except PermissionError:
                        st.error("❌ El archivo de anexo está abierto en otro programa. Ciérralo y vuelve a presionar Guardar.")
                    except Exception as e:
                        st.error(f"❌ Error al guardar el anexo: {e}")


